"""Build an Edme-branded multi-sheet Excel workbook of the lease reports."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EDME_BLUE = "1C3687"
EDME_RED = "ED1B2F"
HEADER_FILL = PatternFill("solid", fgColor=EDME_BLUE)
SUBHEAD_FILL = PatternFill("solid", fgColor="EEF1F8")
WHITE_BOLD = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="FFFFFF", bold=True, size=16)
SUB_FONT = Font(color="FFFFFF", size=10)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D9DEE8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _brand_header(ws, entity_name: str, subtitle: str, ncols: int) -> int:
    last = get_column_letter(max(ncols, 4))
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws.merge_cells(f"A3:{last}3")
    ws["A1"] = "edme  ·  LeasePro AI"
    ws["A1"].font = SUB_FONT
    ws["A2"] = entity_name
    ws["A2"].font = TITLE_FONT
    ws["A3"] = f"{subtitle}   ·   Generated {date.today().strftime('%d %b %Y')}"
    ws["A3"].font = SUB_FONT
    for r in (1, 2, 3):
        ws[f"A{r}"].fill = HEADER_FILL
        ws[f"A{r}"].alignment = Alignment(horizontal="left", vertical="center")
        for c in range(1, max(ncols, 4) + 1):
            ws.cell(row=r, column=c).fill = HEADER_FILL
    ws.row_dimensions[2].height = 26
    return 5  # first data row


def _table(ws, start_row: int, headers: list[str], rows: list[list]):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = WHITE_BOLD
        c.fill = PatternFill("solid", fgColor=EDME_RED)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=start_row + i, column=j, value=val)
            c.border = BORDER
            if isinstance(val, (int, float, Decimal)):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
    for j, h in enumerate(headers, start=1):
        width = max(len(str(h)) + 2, *(len(str(r[j - 1])) + 2 for r in rows)) if rows else len(str(h)) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 12), 40)


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def build_reports_workbook(entity_name: str, data: dict) -> bytes:
    wb = Workbook()

    # Lease Register
    ws = wb.active
    ws.title = "Lease Register"
    r = _brand_header(ws, entity_name, "Lease Register", 7)
    _table(
        ws, r,
        ["Lease Code", "Asset", "Category", "Lessor", "Commencement", "Lease Liability", "ROU NBV"],
        [
            [l.get("lease_code"), l.get("asset_name"), l.get("asset_category"),
             (l.get("lessors") or {}).get("name") if l.get("lessors") else "",
             l.get("commencement_date"), _num(l.get("current_lease_liability")), _num(l.get("current_rou_nbv"))]
            for l in data.get("lease_register", [])
        ],
    )

    # Liability Rollforward (one sheet, all leases stacked)
    ws = wb.create_sheet("Liability Rollforward")
    r = _brand_header(ws, entity_name, "Lease Liability Rollforward", 8)
    rows = []
    for lease in data.get("liability_rollforward", []):
        for s in lease.get("schedule", []):
            rows.append([lease.get("lease_code"), s.get("period_number"), s.get("period_start"), s.get("period_end"),
                         _num(s.get("opening_liability")), _num(s.get("interest_expense")), _num(s.get("payment")), _num(s.get("closing_liability"))])
    _table(ws, r, ["Lease", "Period", "Start", "End", "Opening", "Interest", "Payment", "Closing"], rows)

    # ROU Register
    ws = wb.create_sheet("ROU Register")
    r = _brand_header(ws, entity_name, "Right-of-Use Asset Register", 7)
    rows = []
    for lease in data.get("rou_register", []):
        for s in lease.get("schedule", []):
            rows.append([lease.get("lease_code"), s.get("period_number"), s.get("period_end"),
                         _num(s.get("opening_nbv")), _num(s.get("depreciation")), _num(s.get("impairment")), _num(s.get("closing_nbv"))])
    _table(ws, r, ["Lease", "Period", "End", "Opening NBV", "Depreciation", "Impairment", "Closing NBV"], rows)

    # Security Deposits
    ws = wb.create_sheet("Security Deposits")
    r = _brand_header(ws, entity_name, "Security Deposit Register", 6)
    rows = []
    for lease in data.get("deposit_register", []):
        for d in lease.get("deposits", []):
            rows.append([lease.get("lease_code"), d.get("paid_date"), _num(d.get("deposit_amount")),
                         _num(d.get("present_value")), _num(d.get("prepaid_rent_component")), d.get("expected_refund_date")])
    _table(ws, r, ["Lease", "Paid Date", "Deposit Amount", "Present Value", "Prepaid Rent", "Refund Date"], rows)

    # Disclosures summary
    ws = wb.create_sheet("Disclosures")
    r = _brand_header(ws, entity_name, "Ind AS 116 / IFRS 16 Disclosures", 4)
    disc = data.get("disclosures", {})
    _table(ws, r, ["Measure", "Amount"], [
        ["Total ROU asset (NBV)", _num(disc.get("total_rou_asset_nbv"))],
        ["Total lease liability", _num(disc.get("total_lease_liability"))],
        ["Weighted-average discount rate", _num(disc.get("weighted_average_discount_rate"))],
        ["Total cash outflow (YTD)", _num(disc.get("total_cash_outflow_ytd"))],
    ])
    r2 = r + 7
    ws.cell(row=r2, column=1, value="Maturity analysis (undiscounted)").font = BOLD
    _table(ws, r2 + 1, ["Bucket", "Undiscounted"], [
        [b.get("label"), _num(b.get("undiscounted_amount"))] for b in disc.get("maturity_analysis", [])
    ])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
